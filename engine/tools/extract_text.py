"""Read the user's files: pull the text (and the figures) out of PDF, Word, PowerPoint, Excel, CSV and text files so
Claude can read them in pieces. The user's files are only read, never changed.
  .aura/venv/Scripts/python.exe .aura/engine/tools/extract_text.py ["3 - Put your files here"] [--out .aura/temp/text] [--no-images]
Writes one .txt per file into the out folder (same sub-folders) and the pictures found inside documents into
<out>/<file>.images/. Prints a table: file, kind, pages/sheets, characters, where the text went, and warnings
(e.g. a scanned PDF with no text: look at its page pictures instead)."""
import csv, io, re, sys, zipfile
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

MIN_IMG = 160                     # skip icons and bullets smaller than this (px, both sides)
MAX_ROWS = 400                    # rows per sheet / CSV written out (the total is reported)
TEXT_EXT = {'.txt', '.md', '.markdown', '.rtf', '.tex', '.bib', '.json', '.xml', '.html', '.htm', '.py', '.m', '.ris'}
IMG_EXT = {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff', '.webp', '.emf', '.wmf', '.svg'}


def read_any(p: Path):
    raw = p.read_bytes()
    for enc in ('utf-8-sig', 'utf-16', 'cp1252', 'latin-1'):
        try:
            t = raw.decode(enc)
            if enc == 'utf-16' and not raw[:2] in (b'\xff\xfe', b'\xfe\xff'):
                continue
            return t
        except Exception:
            continue
    return raw.decode('utf-8', 'replace')


def save_image(data: bytes, folder: Path, name: str, saved: list):
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        w, h = im.size
        if w < MIN_IMG or h < MIN_IMG:
            return
        folder.mkdir(parents=True, exist_ok=True)
        ext = (im.format or 'png').lower().replace('jpeg', 'jpg')
        if ext not in ('png', 'jpg', 'gif', 'webp'):
            ext = 'png'
            buf = io.BytesIO()
            im.convert('RGBA' if im.mode in ('RGBA', 'LA', 'P') else 'RGB').save(buf, 'PNG')
            data = buf.getvalue()
        target = folder / f'{name}.{ext}'
        target.write_bytes(data)
        saved.append(f'{target.name} ({w}x{h})')
    except Exception:
        pass


def pdf(p: Path, img_dir, saved):
    from pypdf import PdfReader
    r = PdfReader(str(p))
    if r.is_encrypted:
        try:
            r.decrypt('')
        except Exception:
            return '', 0, ['password protected: ask the user for an unlocked copy']
    out, warn = [], []
    for i, page in enumerate(r.pages, 1):
        try:
            t = page.extract_text() or ''
        except Exception:
            t = ''
        out.append(f'\n--- page {i} ---\n{t.strip()}')
        if img_dir:
            try:
                for k, im in enumerate(page.images, 1):
                    save_image(im.data, img_dir, f'p{i:03d}-{k}', saved)
            except Exception:
                pass
    text = '\n'.join(out)
    if len(re.sub(r'\s|--- page \d+ ---', '', text)) < 40 * max(1, len(r.pages)) * 0.2:
        warn.append('very little text: probably scanned pages; look at the saved page pictures instead')
    return text, len(r.pages), warn


def docx(p: Path, img_dir, saved):
    import docx as D
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    d = D.Document(str(p))
    out = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            para = Paragraph(child, d)
            t = para.text.strip()
            if not t:
                continue
            style = (para.style.name if para.style is not None else '') or ''
            m = re.match(r'Heading (\d)', style)
            out.append(('#' * int(m.group(1)) + ' ' + t) if m else ('# ' + t if style == 'Title' else t))
        elif tag == 'tbl':
            table = Table(child, d)
            out.append('')
            for row in table.rows:
                cells = []
                for c in row.cells:
                    v = c.text.strip().replace('\n', ' ')
                    if not cells or cells[-1] != v:
                        cells.append(v)
                out.append('| ' + ' | '.join(cells) + ' |')
            out.append('')
    if img_dir:
        with zipfile.ZipFile(p) as z:
            for n in z.namelist():
                if n.startswith('word/media/'):
                    save_image(z.read(n), img_dir, Path(n).stem, saved)
    return '\n'.join(out), None, []


def pptx(p: Path, img_dir, saved):
    from pptx import Presentation
    prs = Presentation(str(p))
    out = []
    for i, s in enumerate(prs.slides, 1):
        out.append(f'\n--- slide {i} ---')
        for sh in s.shapes:
            if sh.has_text_frame:
                t = '\n'.join(pg.text for pg in sh.text_frame.paragraphs if pg.text.strip())
                if t.strip():
                    out.append(t.strip())
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows:
                    out.append('| ' + ' | '.join(c.text.strip() for c in row.cells) + ' |')
            if img_dir and sh.shape_type == 13:
                try:
                    save_image(sh.image.blob, img_dir, f's{i:03d}-{sh.shape_id}', saved)
                except Exception:
                    pass
        if s.has_notes_slide:
            n = s.notes_slide.notes_text_frame.text.strip()
            if n:
                out.append('[notes] ' + n)
    return '\n'.join(out), len(prs.slides), []


def xlsx(p: Path, img_dir, saved):
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    out, warn = [], []
    for ws in wb.worksheets:
        rows = 0
        out.append(f'\n--- sheet "{ws.title}" ---')
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None or str(v).strip() == '' for v in row):
                continue
            rows += 1
            if rows <= MAX_ROWS:
                vals = ['' if v is None else (f'{v:.6g}' if isinstance(v, float) else str(v)) for v in row]
                while vals and vals[-1] == '':
                    vals.pop()
                out.append(', '.join(vals))
        if rows > MAX_ROWS:
            out.append(f'... {rows - MAX_ROWS} more rows ({rows} in total)')
            warn.append(f'sheet "{ws.title}" has {rows} rows; first {MAX_ROWS} written')
    n = len(wb.worksheets)
    wb.close()
    return '\n'.join(out), n, warn


def csvfile(p: Path, img_dir, saved):
    text = read_any(p)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
    except Exception:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    out = [', '.join(r) for r in rows[:MAX_ROWS]]
    warn = []
    if len(rows) > MAX_ROWS:
        out.append(f'... {len(rows) - MAX_ROWS} more rows ({len(rows)} in total)')
        warn.append(f'{len(rows)} rows; first {MAX_ROWS} written')
    return '\n'.join(out), None, warn


def textfile(p: Path, img_dir, saved):
    t = read_any(p)
    if p.suffix.lower() == '.rtf':
        t = re.sub(r'\\[a-z]+-?\d* ?|[{}]', '', t)
    return t, None, []


def image(p: Path, img_dir, saved):
    try:
        from PIL import Image
        with Image.open(p) as im:
            return '', None, [f'picture {im.size[0]}x{im.size[1]} px: look at it directly']
    except Exception:
        return '', None, ['picture: look at it directly']


KINDS = {'.pdf': ('PDF', pdf), '.docx': ('Word', docx), '.pptx': ('PowerPoint', pptx), '.xlsx': ('Excel', xlsx),
         '.xlsm': ('Excel', xlsx), '.csv': ('CSV', csvfile), '.tsv': ('CSV', csvfile)}


def main():
    args = sys.argv[1:]
    out_dir, images = Path('.aura/temp/text'), True
    if '--out' in args:
        i = args.index('--out'); out_dir = Path(args[i + 1]); del args[i:i + 2]
    if '--no-images' in args:
        images = False; args.remove('--no-images')
    src = Path(args[0] if args else '3 - Put your files here')
    if not src.exists():
        print(f'Not found: {src}')
        return 2
    base = src if src.is_dir() else src.parent
    files = sorted(f for f in (src.rglob('*') if src.is_dir() else [src])
                   if f.is_file() and not f.name.startswith(('~$', '.')) and f.name.lower() != 'desktop.ini')
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = []
    for f in files:
        ext = f.suffix.lower()
        kind, fn = KINDS.get(ext, (None, None))
        if not fn and ext in TEXT_EXT:
            kind, fn = 'Text', textfile
        if not fn and ext in IMG_EXT:
            kind, fn = 'Picture', image
        rel = f.relative_to(base)
        if not fn:
            label = {'.doc': 'old Word format: ask for a .docx or PDF copy', '.ppt': 'old PowerPoint format: ask for a .pptx copy',
                     '.xls': 'old Excel format: ask for a .xlsx or CSV copy'}.get(ext, 'not read (unknown kind)')
            rows.append((str(rel), ext.lstrip('.') or '?', '', 0, '', [label]))
            continue
        target = out_dir / rel.with_name(rel.name + '.txt')
        img_dir = (out_dir / rel.with_name(rel.name + '.images')) if images and kind not in ('Picture', 'Text', 'CSV', 'Excel') else None
        saved = []
        try:
            text, count, warn = fn(f, img_dir, saved)
        except Exception as e:
            rows.append((str(rel), kind, '', 0, '', [f'could not read: {str(e)[:90]}']))
            continue
        where = ''
        if text.strip():
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(f'# {rel.as_posix()}\n{text.strip()}\n', encoding='utf-8')
            where = str(target).replace('\\', '/')
        if saved:
            warn = warn + [f'{len(saved)} picture(s) saved to {str(img_dir).replace(chr(92), "/")}/']
        rows.append((str(rel).replace('\\', '/'), kind, count or '', len(text), where, warn))
    if not rows:
        print(f'No files in {src}.')
        return 0
    print(f'Read {len(rows)} file(s) from {str(src).replace(chr(92), "/")}:')
    for name, kind, count, chars, where, warn in rows:
        unit = {'PDF': 'page', 'PowerPoint': 'slide', 'Excel': 'sheet'}.get(kind, '') + ('s' if count != 1 else '')
        line = f'  - {name} [{kind}' + (f', {count} {unit}' if count else '') + (f', {chars:,} chars' if chars else '') + ']'
        if where:
            line += f' -> {where}'
        print(line)
        for w in warn:
            print(f'      note: {w}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
