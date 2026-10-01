"""Speaker notes (and a timed script) as a Word file: a small picture of each slide, its time window and its notes.
  .aura/venv/Scripts/python.exe .aura/engine/tools/export_notes.py "4 - Your slides/<Title>.html" [--timed] ["<out.docx>"]
Without --timed the time windows are left out. Time windows come from each slide's data-minutes."""
import sys
from pathlib import Path

sys.dont_write_bytecode = True          # keep the engine folder clean
sys.path.insert(0, str(Path(__file__).resolve().parent))
from export_pptx import render, aura_root  # noqa: E402

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass


def mmss(minutes):
    s = round(minutes * 60)
    return f'{s // 60}:{s % 60:02d}'


def main():
    args = sys.argv[1:]
    timed = '--timed' in args
    args = [a for a in args if a != '--timed']
    if not args:
        print(__doc__)
        return 2
    deck = Path(args[0]).resolve()
    if deck.is_dir():
        deck = deck / 'index.html'
    if not deck.is_file():
        print(f'Deck not found: {deck}')
        return 2
    out = Path(args[1]).resolve() if len(args) > 1 else deck.with_name(deck.stem + ' - speaker notes.docx')
    folder, meta = render(deck)

    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    doc = Document()
    st = doc.styles['Normal']
    st.font.name, st.font.size = 'Calibri', Pt(13)
    title = meta.get('title') or deck.stem
    doc.add_heading(title, 0)
    total = sum(s.get('minutes') or 0 for s in meta['slides'])
    intro = f"{len(meta['slides'])} slides"
    if timed and total:
        intro += f', about {mmss(total)} minutes in total'
    doc.add_paragraph(intro + '. Read the notes in your own words; they are reminders, not a script to read aloud.' if not timed
                      else intro + '. The times show where you should be; a little early or late is fine.')
    elapsed = 0.0
    for s in meta['slides']:
        head = f"Slide {s['number']}" + (f" - {s['title']}" if s.get('title') else '')
        doc.add_heading(head, level=2)
        if timed and s.get('minutes'):
            p = doc.add_paragraph()
            r = p.add_run(f"{mmss(elapsed)} to {mmss(elapsed + s['minutes'])}  ({s['minutes']:g} min)")
            r.bold = True
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            elapsed += s['minutes']
        doc.add_picture(str(folder / s['image']), width=Inches(3.6))
        notes = (s.get('notes') or '').strip()
        for para in (notes.split('\n\n') if notes else ['(no notes for this slide)']):
            doc.add_paragraph(para.strip())
    tmp = out.with_name(out.name + '.part')
    doc.save(tmp)
    tmp.replace(out)
    root = aura_root(deck)
    shown = str(out.relative_to(root) if root and str(out).startswith(str(root)) else out).replace('\\', '/')
    with_notes = sum(1 for s in meta['slides'] if (s.get('notes') or '').strip())
    print(f'Speaker notes saved: {shown}')
    print(f"  {len(meta['slides'])} slides, notes on {with_notes}" + (f', timed script {mmss(total)} min' if timed and total else ''))
    return 0


if __name__ == '__main__':
    sys.exit(main())
